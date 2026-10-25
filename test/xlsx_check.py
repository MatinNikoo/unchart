"""Reads the native chart out of an Unchart XLSX and prints what it holds as JSON, for test/editor.js."""
import json
import sys
import zipfile

import openpyxl


def text_of(t):
    try:
        return ''.join(r.t for p in t.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return None


def rgb(fill):
    """The hex color of a solid fill, whatever openpyxl wraps it in."""
    if fill is None or fill.srgbClr is None:
        return None
    c = fill.srgbClr
    return getattr(c, 'val', c)


def color_of(gp):
    return rgb(gp.solidFill)


path = sys.argv[1]
ws = openpyxl.load_workbook(path)['Unchart']
out = {'rows': ws.max_row, 'a1': ws['A1'].value, 'charts': [], 'parts': []}
with zipfile.ZipFile(path) as z:
    out['parts'] = sorted(z.namelist())
    drawing = [n for n in z.namelist() if n.startswith('xl/drawings/drawing') and n.endswith('.xml')]
    out['descr'] = any('descr="' in z.read(n).decode('utf8') for n in drawing)
    out['userShapes'] = [n for n in z.namelist() if 'userShapes' in n or 'drawings/drawing' in n and 'chart' in n]
    chart_xml = ''.join(z.read(n).decode('utf8') for n in z.namelist() if n.startswith('xl/charts/chart'))
    out['xml'] = {
        'gradFill': '<a:gradFill' in chart_xml,
        'pattFill': '<a:pattFill' in chart_xml,
        'prstDash': '<a:prstDash' in chart_xml,
        'effectLst': '<a:effectLst' in chart_xml,
        'errBars': '<c:errBars' in chart_xml,
        'dLbls': '<c:dLbls' in chart_xml,
        'userShapes': '<c:userShapes' in chart_xml,
        'smooth1': '<c:smooth val="1"' in chart_xml,
    }
for c in ws._charts:
    d = {'type': type(c).__name__, 'series': sum(len(q.series) for q in getattr(c, '_charts', [c])), 'plots': [type(q).__name__ for q in getattr(c, '_charts', [c])], 'title': text_of(c.title) if c.title else None,
         'legend': c.legend.position if c.legend else None, 'grouping': getattr(c, 'grouping', None),
         'gapWidth': getattr(c, 'gapWidth', None), 'overlap': getattr(c, 'overlap', None), 'barDir': getattr(c, 'barDir', None)}
    s0 = c.series[0] if c.series else None
    if s0 is not None:
        gp = s0.graphicalProperties
        d['s0'] = {'fill': color_of(gp) if gp is not None else None,
                   'line': rgb(gp.line.solidFill) if gp is not None and gp.line is not None else None,
                   'marker': s0.marker.symbol if getattr(s0, 'marker', None) else None,
                   'smooth': bool(getattr(s0, 'smooth', False))}
    for nm in ('x_axis', 'y_axis'):
        ax = getattr(c, nm, None)
        if ax is not None:
            d[nm] = {'min': ax.scaling.min, 'max': ax.scaling.max, 'orientation': ax.scaling.orientation, 'log': ax.scaling.logBase,
                     'major': getattr(ax, 'majorUnit', None), 'numFmt': ax.numFmt.formatCode if ax.numFmt is not None else None,
                     'title': text_of(ax.title) if ax.title else None, 'delete': ax.delete, 'majorGrid': ax.majorGridlines is not None}
    out['charts'].append(d)
print(json.dumps(out))
